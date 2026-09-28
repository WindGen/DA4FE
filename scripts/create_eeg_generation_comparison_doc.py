from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"E:\TeCh-improve\EEG图片生成模型对比与DA4FE适配方案.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def format_paragraph(p, space_after=6, line=1.18, first_indent=True):
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if first_indent:
        pf.first_line_indent = Pt(21)


def add_para(doc, text="", bold_prefix=None, space_after=6, first_indent=True):
    p = doc.add_paragraph()
    format_paragraph(p, space_after=space_after, first_indent=first_indent)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    pf = p.paragraph_format
    pf.space_after = Pt(3)
    pf.line_spacing = 1.12
    r = p.add_run(text)
    set_run_font(r, size=10.2)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, size=14 if level == 1 else 11.5, bold=True)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(h))
        set_run_font(r, size=font_size, bold=True, color="FFFFFF")
        set_cell_shading(hdr[i], "1F4E78")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(hdr[i])
        set_cell_margins(hdr[i])
        if widths:
            hdr[i].width = Inches(widths[i])
    mark_header_row(table.rows[0])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2 == 1:
                set_cell_shading(cells[i], "F2F6FA")
            set_cell_border(cells[i])
            set_cell_margins(cells[i])
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_source_table(doc, rows):
    add_table(doc, ["项目", "源码位置", "核对内容"], rows, widths=[1.35, 3.0, 2.35], font_size=8.6)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("DA4FE EEG 图片生成技术分析")
    set_run_font(footer_run, size=8.5, color="666666")
    return doc


def build_doc():
    doc = setup_document()

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(32)
    title.paragraph_format.space_after = Pt(12)
    r = title.add_run("EEG 图片生成模型对比与 DA4FE 适配方案")
    set_run_font(r, size=22, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    r = subtitle.add_run("基于 DreamDiffusion main 与 AVDE master 源码的技术说明")
    set_run_font(r, size=11, color="555555")

    add_para(doc, "本文分析 E Model & Code 目录中的 DreamDiffusion-main.zip 和 avde-master.zip，重点说明 EEG 如何进入图像生成模型、EEG embedding 在训练中的作用，以及这两种路线与当前 DA4FE Stage-1 和 Stage-2 的衔接方式。结论先行：在当前 DA4FE 已能稳定输出 EEG 特征、而数据管线仍以 EEG 和类别字段为主的条件下，优先采用 DreamDiffusion 的扩散条件适配路线；AVDE 更适合在拥有大量严格配对的 EEG 图片数据、较高显存预算和长期追求高吞吐自回归生成时使用。", space_after=10, first_indent=False)

    add_heading(doc, "一 执行摘要", 1)
    add_para(doc, "DreamDiffusion 的核心是把 EEG encoder 输出映射为 Stable Diffusion U-Net 的 cross-attention 条件。它先用掩码自编码器预训练 EEG 表征，再冻结图像 VAE，主要训练 EEG 条件分支、条件投影和部分 U-Net 条件注意力，并叠加图像扩散噪声损失与 CLIP 图像 embedding 对齐损失。其改造点集中在条件侧，和 DA4FE 的现有特征接口最接近。", space_after=5)
    add_para(doc, "AVDE 的核心是先用 LaBraM 将 EEG 映射到 OpenCLIP 图像 embedding 空间，再用冻结的 VQ-VAE 把图片离散化成视觉 token，训练 VAREEG 视觉自回归 Transformer 按多尺度顺序预测这些 token。它不是像素回归，而是 EEG 条件下的离散视觉 token 生成，工程链路更长、显存和数据要求更高。", space_after=5)
    add_para(doc, "对 DA4FE 的推荐落地顺序是：保留 DA4FE Stage-1 作为 EEG encoder，新增一个小型 token projector；第一版冻结扩散模型的 VAE 和大部分 U-Net，只训练 projector、条件映射和 cross-attention；确认图像语义指标提升后，再考虑解冻更多 U-Net 或联合微调 DA4FE。", space_after=8)

    add_heading(doc, "二 分析范围和当前 DA4FE 基线", 1)
    add_para(doc, "源码分析对象为两个压缩包解压后的代码，不依据项目名称推断未实现的功能。源码位置统一以压缩包根目录下的相对路径表示。当前 DA4FE 代码位于本工作区 E TeCh-improve。", space_after=5)
    add_table(doc, ["项目", "当前代码行为", "对图像生成的含义"], [
        ("DA4FE Stage-1", "通道、时间、频率三分支编码，BranchFusion 得到 fused feature，再经 projector 输出类别 logits；支持 CE、Triplet、Multi-Similarity、ArcFace、CosFace 等损失。", "已有可复用的 EEG 表征，但 projector 的类别 logits 不是图像生成条件，生成阶段应使用 fused feature。"),
        ("Stage-1 特征输出", "Stage1FeatureModel 可对 fused feature 做最终 L2 normalize；默认参数中 stage1_l2_normalize=True。", "可作为稳定的条件向量输入，但生成 adapter 仍需要学习尺度和 token 结构。"),
        ("Stage-2", "当前是基于 Stage-1 特征的 sklearn 分类器集合，不是图像解码器。", "不能直接生成图片；可保留作语义一致性评估或辅助分类器。"),
        ("当前 HF loader", "主要保留 conditioning_image、label、subject、label_folder，其他字段会被删除。", "若要训练生成器，必须额外恢复每个 EEG 样本对应的原始图片路径或图片像素。"),
    ], widths=[1.25, 3.35, 2.1], font_size=8.6)
    add_para(doc, "当前 DA4FE 默认输入配置为 128 个 EEG 通道、目标序列长度 512，并启用通道、时间和频率分支。默认融合输出维度为 512，最终分类类别数由数据集解析得到。生成阶段应从 `backbone(..., return_features=True)` 取得 fused feature，而不是把 Stage-1 的类别 logits 当作 EEG embedding。", space_after=6)

    add_heading(doc, "三 DreamDiffusion 实现方式", 1)
    add_heading(doc, "三点一 EEG 预训练", 2)
    add_para(doc, "DreamDiffusion 的 Stage A 使用 `sc_mbm.MAEforEEG`。输入 EEG 先经过一维 patch embedding，再用 Transformer encoder 处理；训练时随机遮挡 patch，只对被遮挡 patch 做重建损失。该阶段的目标是让 encoder 学到对噪声和个体差异更稳健的 EEG 表征，而不是直接生成图片。代码还提供可选的自然图像重建特征损失，但默认配置中 `use_nature_img_loss=False`。", space_after=5)
    add_para(doc, "预训练 encoder 的典型输出是形状 `[B, N, 1024]` 的 EEG token 序列。`global_pool=True` 时会平均为一个 token；`global_pool=False` 时保留 patch 序列，随后可以用通道映射得到 77 个条件 token。", space_after=6)

    add_heading(doc, "三点二 EEG 条件如何进入 Stable Diffusion", 2)
    add_table(doc, ["步骤", "张量或模块", "作用"], [
        ("一", "EEG encoder", "读取 EEG，输出约 1024 维的 patch token 或全局 token。"),
        ("二", "channel_mapper 可选", "global_pool=False 时把 patch 序列长度映射到 77 个 token。"),
        ("三", "dim_mapper", "把 EEG token 的最后一维从 1024 映射到 Stable Diffusion 的 context_dim 768。"),
        ("四", "U-Net cross attention", "把 EEG 条件作为 cross-attention context，影响每个扩散时间步的噪声预测。"),
        ("五", "first stage VAE", "将真实图片编码成 4 通道 latent，扩散模型在 latent 空间预测噪声，最后用 VAE 解码图片。"),
    ], widths=[0.55, 1.8, 4.35], font_size=8.8)
    add_para(doc, "因此，DreamDiffusion 不是把 EEG 直接拼到图片像素上，而是把 EEG 变成扩散 U-Net 可以读取的条件 token。Stable Diffusion 的图像 VAE 和大部分生成先验来自预训练模型，EEG 分支负责把脑信号转成与图像语义和去噪过程兼容的条件。", space_after=6)

    add_heading(doc, "三点三 训练目标和冻结策略", 2)
    add_para(doc, "每个训练样本包含 EEG、类别和对应图片。扩散部分随机采样时间步，把图片 latent 加噪，U-Net 预测噪声，产生标准 latent diffusion loss。若 `clip_tune=True`，代码用冻结的 CLIP 图像 encoder 得到真实图片 embedding，再用一个 1024 到 768 的 mapping 对 EEG encoder latent 做投影，加入 `1 - cosine_similarity` 的图像语义对齐损失；若打开 `cls_tune`，还可以加入 40 类交叉熵。默认配置开启 CLIP 对齐、关闭分类微调。", space_after=5)
    add_para(doc, "源码的 `finetune` 会先解冻整体模型再冻结 first stage。随后把 `train_cond_stage_only=True`，优化器实际包含 EEG 条件分支参数，以及名字包含 `attn2`、`time_embed_condtion` 和 `norm2` 的条件相关 U-Net 参数。因此，默认并非只训练一个线性投影器，也不是完整解冻 Stable Diffusion；它是在冻结 VAE 的前提下，联合微调 EEG 条件侧和部分 U-Net 条件层。", space_after=6)

    add_heading(doc, "三点四 推理流程", 2)
    add_para(doc, "推理时把 EEG 经过条件模型得到 context，使用 PLMS 采样器从随机 latent 开始迭代去噪，得到 4 通道图像 latent，再调用 Stable Diffusion 的 first-stage decoder 输出图片。一个 EEG 可以采样多张图片，因此输出具有生成多样性；评估代码计算 MSE、PCC、SSIM、PSM 和基于类别的 top-1 指标。", space_after=6)

    add_heading(doc, "四 AVDE 实现方式", 1)
    add_heading(doc, "四点一 数据和 EEG CLIP 对齐", 2)
    add_para(doc, "AVDE 的 `EEGDataset` 读取 Things-EEG 预处理结果。训练集按图像条件组织，每个类别包含多次 EEG 重复；数据集同时返回 EEG、图片像素、类别标签、文本描述，以及预先计算的 OpenCLIP ViT-H-14 文本和图像 embedding。图片 embedding 维度为 1024，与 LaBraM 的输出 head 对齐。", space_after=5)
    add_para(doc, "第一阶段使用 `NeuralTransformer` 作为 LaBraM EEG encoder。它先做时间卷积或 patch embedding，再加入通道和时间位置编码，经过 Transformer block 和平均池化，最后通过 `head` 输出 1024 维 EEG feature。`engine_for_finetuning.py` 用该 feature 和对应图片 embedding 计算 MSE regression loss 与 CLIP 对比损失的加权和，代码中 alpha=0.8，即回归项权重更大。虽然模型 forward 同时计算 text loss，当前训练循环的总损失只使用回归项和 image CLIP loss，没有把 text loss 加入总损失。", space_after=6)

    add_heading(doc, "四点二 图像 tokenizer 和视觉自回归生成", 2)
    add_para(doc, "第二阶段加载预训练 VQ-VAE，把真实图片编码并量化成离散视觉 token。`VAREEG` 的 `embed_proj` 将 1024 维 EEG embedding 投影到自回归 Transformer 的条件维度；该条件通过 AdaLN 调制每个 Transformer block。模型按多尺度 patch 序列预测下一个视觉 token，训练损失是离散 token 的交叉熵，而不是像素 MSE。", space_after=5)
    add_para(doc, "推理时，模型从 EEG 条件生成第一层视觉 token，随后逐尺度生成更高分辨率 token，支持 classifier-free guidance、top-k 和 top-p 采样，最后由 VQ-VAE 解码成图片。代码中的 `label` 不作为 VAR 的类别条件输入，生成的主要条件是 EEG embedding；类别标签主要参与数据索引和第一阶段的检索式评估。", space_after=6)

    add_heading(doc, "四点三 训练冻结关系", 2)
    add_para(doc, "`train_eeg.py` 加载预训练 VQ-VAE、VAR checkpoint 和微调后的 LaBraM。默认 `unfreeze_labram=False`，LaBraM 参数被冻结；优化器只从 VAREEG 参数建立参数组，因此第二阶段主要训练视觉自回归生成器。VQ-VAE 在构建时以 test mode 加载，作为固定图像 tokenizer。", space_after=6)

    add_heading(doc, "五 两种方法的关键差异", 1)
    add_table(doc, ["维度", "DreamDiffusion", "AVDE", "对 DA4FE 的影响"], [
        ("生成空间", "Stable Diffusion 连续 latent", "VQ-VAE 离散视觉 token", "连续扩散适配更直接；离散 token 需要新增 tokenizer 和 AR 训练。"),
        ("EEG 条件注入", "U-Net cross-attention context", "VAR 的 embed_proj 和 AdaLN 条件", "DA4FE 可直接接 projector；两者都不需要把 DA4FE 改成分类器。"),
        ("EEG 预训练", "掩码 EEG 自编码器", "LaBraM + CLIP 图像对齐", "DA4FE 已有 Stage-1 表征，可跳过两套 EEG backbone 的重复训练。"),
        ("图像监督", "扩散噪声损失加 CLIP 图像语义损失", "图片 VQ token 交叉熵，另有 CLIP 对齐预训练", "两者都要求 EEG 和原始图片一一配对；只有类别标签不够。"),
        ("默认微调范围", "冻结 VAE，训练条件侧和部分 U-Net 条件层", "VQ-VAE 和 LaBraM 默认冻结，训练 VAR", "DreamDiffusion 初始显存和工程改造成本更低。"),
        ("采样特点", "多步扩散，质量和多样性较好，速度较慢", "多尺度自回归，token 采样速度和控制方式不同", "先做小规模扩散验证更容易定位问题。"),
        ("数据规模要求", "可在较小配对集上依赖 SD 先验，但仍易过拟合", "对图像 token 分布和 EEG 图像配对更敏感", "当前 40 类数据优先使用冻结生成先验，降低过拟合风险。"),
    ], widths=[1.25, 2.05, 2.05, 1.45], font_size=8.0)

    add_heading(doc, "六 与当前 DA4FE 的适配判断", 1)
    add_heading(doc, "六点一 推荐 DreamDiffusion 路线的原因", 2)
    add_bullet(doc, "接口匹配。DA4FE Stage-1 已能输出固定维度 fused feature；DreamDiffusion 的 cond_stage_model 本质上需要一个 EEG encoder 输出和一个维度映射，适配范围集中在 projector 和 token 化。")
    add_bullet(doc, "保留已有训练资产。可以冻结或低学习率微调 DA4FE，不必重写三分支通道、时间和频率编码器，也不必先复刻 LaBraM。")
    add_bullet(doc, "图像先验更强。Stable Diffusion 直接提供连续图像生成先验，第一版只训练 EEG 条件相关参数，通常比从 EEG 直接学习完整视觉 token 分布更稳。")
    add_bullet(doc, "适合逐步验证。可以先只做 40 类小规模实验，使用 CLIP image similarity、类别检索和人工可辨识度检查，确认 EEG 是否真正影响图片，再逐步放开 U-Net。")

    add_heading(doc, "六点二 AVDE 适用条件", 2)
    add_para(doc, "AVDE 不是不可用，而是更适合作为第二条实验路线：当你能够提供稳定的 EEG 图片配对、较大的训练批量、VQ-VAE 和 VAR 预训练权重，并希望研究离散视觉 token 的长程生成时再采用。将 DA4FE 接入 AVDE 时，需要把 DA4FE feature 投影到 1024 维，替换 `labram_local.forward` 的输出，同时改写数据和训练脚本；如果只保留 40 类，仍需重新确认 VQ-VAE 和 VAR 的图像分布是否适合这套数据。", space_after=6)

    add_heading(doc, "七 DA4FE 的建议落地方案", 1)
    add_heading(doc, "七点一 数据配对先行", 2)
    add_para(doc, "当前 HF loader 在 `_prepare_dataset` 中只保留 EEG、标签、被试和类别文件夹字段。生成训练必须为每条 EEG 恢复对应图片路径或图片 tensor，建议新增 `image_id` 和 `image_path` 字段，并保证同一条 EEG 的图片不会因排序或 split 重排而错配。若使用旧的 EEG-ImageNet PTH 数据，样本已有 `image` 索引，可以沿用 `images[image]` 建立路径。", space_after=6)
    add_heading(doc, "七点二 第一版扩散适配器", 2)
    add_table(doc, ["模块", "建议实现", "首版是否训练"], [
        ("DA4FE encoder", "读取 EEG，输出 fused feature；保留 Stage-1 的归一化开关。", "冻结，先验证条件侧。"),
        ("Feature adapter", "Linear 或两层 MLP，将 DA4FE feature 映射到 768 维；加入 LayerNorm 和 dropout。", "训练。"),
        ("Token projector", "先输出 1 到 4 个 context token；稳定后再扩展到 77 个 token。", "训练。"),
        ("Stable Diffusion VAE", "编码图片 latent 和解码采样结果。", "冻结。"),
        ("U-Net cross attention", "只开放与 context 交互的 cross-attention 参数，必要时再开放条件归一化层。", "训练。"),
        ("语义损失", "真实图片 CLIP embedding 与 EEG adapter 输出做 cosine 对齐。", "训练目标之一。"),
    ], widths=[1.35, 3.85, 1.6], font_size=8.6)
    add_para(doc, "如果使用 1 个 token，工程最简单，适合检查 DA4FE embedding 是否携带可生成的类别和语义信息；如果 1 个 token 的细节不足，再增加可学习 token 化模块。不要第一步就解冻完整 U-Net，否则小数据集很容易学习到图片背景或类别偏差，而不是 EEG 条件。", space_after=6)

    add_heading(doc, "七点三 损失和训练顺序", 2)
    add_bullet(doc, "先训练 projector 和 cross-attention，使用 latent diffusion loss 加较小权重的 CLIP image alignment loss。建议先让扩散损失占主导，再调节 CLIP 权重。")
    add_bullet(doc, "确认 train、validation、test 都按被试隔离，且图片实例或重复 EEG 不跨集合泄漏。当前 DA4FE 已支持 HF 原始划分和按 subject 自划分，应在生成阶段沿用同一 split。")
    add_bullet(doc, "保存 best checkpoint 时不要使用测试集调参。验证集可用于选择 checkpoint，测试集只在最终一次生成评估时使用。")
    add_bullet(doc, "保留 Stage-2 分类器作为独立语义检查：对生成图片做图像分类或 CLIP 检索，不能把 Stage-2 的 EEG 分类准确率直接当作生成质量。")

    add_heading(doc, "八 风险和评价指标", 1)
    add_table(doc, ["风险", "表现", "控制方法"], [
        ("图片错配", "loss 不降或生成图片与类别无关", "按样本 image_id 建立可复现映射，并随机抽样核对 EEG、类别和图片。"),
        ("被试泄漏", "测试结果异常高，换被试性能骤降", "按 subject 划分；自划分固定 seed，并保存 split summary。"),
        ("条件塌缩", "不同 EEG 输出几乎相同图片", "检查 adapter 输出余弦相似度、按类别检索和不同 EEG 的生成差异。"),
        ("生成器过拟合", "训练图片相似，测试图片失真", "冻结更多 SD 参数、增加 dropout、降低学习率和 CLIP 权重，并用验证集选模型。"),
        ("指标误读", "只看 FID 或分类准确率得出错误结论", "同时报告 CLIP image similarity、EEG 到图片 top-k 检索、类别一致性、多样性和人工盲评。"),
    ], widths=[1.35, 2.35, 3.1], font_size=8.6)
    add_para(doc, "建议的最小评估集包括：同一 EEG 的多次采样图、同一类别不同 EEG 的语义一致性、不同类别之间的可分性、跨被试测试，以及生成图片和真实图片的 CLIP 相似度。若后续追求视觉逼真度，再加入 FID 或 KID，但不能只依赖单一指标。", space_after=6)

    add_heading(doc, "九 最终建议", 1)
    add_para(doc, "当前项目优先选择 DreamDiffusion 风格的扩散适配方案。实施上先把 DA4FE Stage-1 的 fused feature 接入一个 768 维 adapter，训练少量 context token 和 U-Net cross-attention；冻结 Stable Diffusion VAE 和主体 U-Net，使用配对图片的 latent diffusion loss 加 CLIP 语义对齐。该方案能最大限度复用 DA4FE 现有模型和训练结果，便于逐层排错。", space_after=6)
    add_para(doc, "AVDE 路线可作为后续对照实验：用 DA4FE 替换 LaBraM，保留 VQ-VAE 和 VAREEG，比较离散 token 生成与扩散生成在类别一致性、跨被试泛化、生成速度和多样性上的差异。除非已有足够图像配对和算力，不建议把 AVDE 作为第一版工程实现。", space_after=8)

    add_heading(doc, "十 源码核对清单", 1)
    add_source_table(doc, [
        ("DreamDiffusion 数据", "DreamDiffusion-main/code/dataset.py:238-341", "EEGDataset 返回 eeg、label、image、image_raw；Splitter 读取 train/test 索引。"),
        ("DreamDiffusion EEG 预训练", "DreamDiffusion-main/code/stageA1_eeg_pretrain.py:123-182", "构建 MAEforEEG、掩码重建训练和 checkpoint 保存。"),
        ("DreamDiffusion 条件模型", "DreamDiffusion-main/code/dc_ldm/ldm_for_eeg.py:29-89", "EEG encoder、可选 channel_mapper、1024 到 context_dim 的 dim_mapper、CLIP mapping。"),
        ("DreamDiffusion 微调", "DreamDiffusion-main/code/dc_ldm/ldm_for_eeg.py:135-169", "冻结 first stage，开启 train_cond_stage_only 并训练条件侧。"),
        ("DreamDiffusion 损失", "DreamDiffusion-main/code/dc_ldm/models/diffusion/ddpm.py:1043-1088, 1160-1198", "扩散噪声损失、CLIP cosine loss 和可选分类损失。"),
        ("AVDE 数据", "avde-master/eeg_data/dataset.py:19-155, 158-355", "加载 EEG、图片、CLIP 文本和图片 embedding，处理训练与测试索引。"),
        ("AVDE LaBraM", "avde-master/models/labram.py:314-458", "EEG Transformer 输出 1024 维 feature，并计算 image/text ClipLoss。"),
        ("AVDE 对齐训练", "avde-master/engine_for_finetuning.py:31-104", "MSE regression 与 image ClipLoss 加权训练；当前 text_loss 未并入总损失。"),
        ("AVDE VAREEG", "avde-master/models/var_eeg.py:21-240", "1024 维 EEG 条件投影、AdaLN 调制、多尺度视觉 token 自回归生成。"),
        ("AVDE 生成训练", "avde-master/train_eeg.py:96-155; avde-master/trainer_eeg.py:58-141", "加载 VQ-VAE、VAR、LaBraM；默认冻结 LaBraM，训练 VAR token loss。"),
        ("当前 DA4FE", "models/DA4FE.py:14-123; exp/exp_stage1_feature.py:87-121", "三分支融合、fused feature、类别 projector 和 Stage-1 L2 normalize。"),
        ("当前 HF 数据管线", "data_provider/data_loader.py:647-715, 1041-1062", "按原始或自定义 subject split 加载，并仅保留 EEG、标签、被试和类别文件夹字段。"),
    ])

    add_para(doc, "文档结论基于上述源码路径中的实际实现。若后续替换数据集、Stable Diffusion 版本或 VQ-VAE 版本，应重新核对 context_dim、embedding 维度、图片配对关系和冻结参数。", space_after=4, first_indent=False)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
