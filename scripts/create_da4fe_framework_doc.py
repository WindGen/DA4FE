from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"E:\TeCh-improve\DA4FE模型框架图绘制说明.docx")

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
LIGHT_ORANGE = "FFF4E5"
LIGHT_GREEN = "EEF8F1"
WHITE = "FFFFFF"
ORANGE = "C55A11"
GREEN = "287D3C"


def set_font(run, name="Microsoft YaHei", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, name="Microsoft YaHei", size=11, color="222222", bold=False):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    total = sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color="D9E2EC", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def shade_paragraph(paragraph, fill=LIGHT_GRAY, left_border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if left_border:
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)
        left = p_bdr.find(qn("w:left"))
        if left is None:
            left = OxmlElement("w:left")
            p_bdr.append(left)
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), left_border)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_font(run, size=9, color=MUTED)


def add_body(doc, text="", bold_prefix=None, color=None, style="Normal"):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_font(r1, size=10.5, color=color or "222222", bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2, size=10.5, color=color or "222222")
    else:
        r = p.add_run(text)
        set_font(r, size=10.5, color=color or "222222")
    return p


def add_flow(doc, text, fill=LIGHT_GRAY, accent=BLUE, font_size=9.5):
    p = doc.add_paragraph(style="Flow")
    shade_paragraph(p, fill=fill, left_border=accent)
    r = p.add_run(text)
    set_font(r, name="Consolas", size=font_size, color=NAVY)
    return p


def add_note(doc, label, text, fill=LIGHT_BLUE, accent=BLUE):
    p = doc.add_paragraph(style="Note")
    shade_paragraph(p, fill=fill, left_border=accent)
    r1 = p.add_run(label + "  ")
    set_font(r1, size=10.5, color=accent, bold=True)
    r2 = p.add_run(text)
    set_font(r2, size=10.5, color="263238")
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_font(r, size={1: 16, 2: 13, 3: 11.5}[level], color=BLUE if level < 3 else DARK_BLUE, bold=True)
    return p


def add_param_table(doc):
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [2300, 2400, 4660])
    set_table_borders(table)
    headers = ["项目", "当前配置", "图中建议标注"]
    for cell, value in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        set_font(r, size=9.5, color=NAVY, bold=True)
    rows = [
        ("输入", "T=500, C=128", "X ∈ R[B,T,C] = [B,500,128]"),
        ("三分支维度", "Dv=Dt=Df=64", "三条分支各输出 64 维向量"),
        ("编码层", "v/t/f = 4/4/4", "CoTAR Encoder ×4"),
        ("patch 长度", "patch_len=4", "Temporal/Frequency Patching"),
        ("频率处理", "fs=1000 Hz, Hann, relative", "251 bins, 0-500 Hz, 2 Hz resolution"),
        ("分支融合", "concat_mlp", "64+64+64=192 → 512 → 256"),
        ("输出", "num_class=40", "[B,256] → Linear → [B,40]"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            r = p.add_run(value)
            set_font(r, size=9, color="263238", bold=(idx == 0))
            if idx == 0:
                set_cell_shading(cells[idx], "F8FAFC")
    return table


def add_prompt_block(doc):
    add_heading(doc, "8. 可直接复制给 AI 绘图软件的完整提示词", 1)
    intro = doc.add_paragraph(style="Normal")
    r = intro.add_run("以下内容可以整体复制。请绘制一张用于深度学习论文的 DA4FE 模型框架图，采用白色背景、横向左到右布局、清晰的矢量风格和学术论文风格。")
    set_font(r, size=10.5, color="222222")
    p = doc.add_paragraph(style="Prompt")
    shade_paragraph(p, fill="F8FAFC", left_border=NAVY)
    prompt = (
        "标题：DA4FE: Dual-Axis and Frequency-enhanced EEG Representation Learning Framework。\n\n"
        "左侧绘制 EEG 输入 X ∈ R[B,T,C]，外部输入格式为 [B,T,C]，当前示例为 [B,500,128]。输入先经过固定长度重采样和可选的逐样本逐通道标准化，然后复制为三条并行分支。三条分支入口分别标注 training-time independent augmentation，可包含 Flip、Channel Mask、Temporal Mask、Frequency Mask、Jitter；三条分支独立选择增强方式。\n\n"
        "顶部蓝色 Channel Branch：输入 X [B,T,C] → Transpose [B,C,T] → augmentation → sinusoidal positional encoding → channel-wise linear embedding。将每一个 EEG 通道视为一个 token，线性层将每个通道的完整时间序列 R^T 映射到 Dv=64 维，得到 channel tokens [B,C,64]，当前为 [B,128,64]。之后进入 CoTAR Encoder ×4，最后沿通道 token 维度进行 Mean Pooling，得到 channel feature zc [B,64]。\n\n"
        "中间绿色 Temporal Branch：输入 X [B,T,C] → Transpose [B,C,T] → augmentation → positional encoding → Cross-Channel Temporal Patching。绘制 Conv2D，kernel=[C,patch_len]，stride=patch_len，当前 patch_len=4。每个时间 patch 同时融合全部 EEG 通道，生成 temporal tokens [B,Nt,64]，当前约为 [B,126,64]。之后进入 CoTAR Encoder ×4，沿时间 token 维度 Mean Pooling，得到 temporal feature zt [B,64]。\n\n"
        "底部橙色或紫色 Frequency Branch：输入 X [B,T,C] → Transpose [B,C,T] → augmentation → channel-wise mean removal → Hann window → one-sided rFFT along time dimension → power spectrum |FFT(x)|² → periodogram PSD normalization → one-sided spectrum correction → relative power normalization → logarithmic power spectrum log(PSD+ε)。当前采样率 fs=1000 Hz，T=500，因此频率分辨率为 2 Hz，频率范围为 0-500 Hz，得到 251 个单边频率 bins。\n\n"
        "频率特征随后经过 Cross-Channel Frequency Patching，kernel=[C,patch_len]，stride=patch_len，得到 frequency tokens [B,Nf,64]，当前约为 [B,63,64]。在频率 token 上增加 Physical Frequency Coordinate Encoding。每个 token 使用真实物理频率 f 的两个坐标：f/fNyquist 和 log(1+f)/log(1+fNyquist)，经过 Linear(2→64) → GELU → Linear(64→64)，再与 frequency tokens 相加。标注 Physical frequency in Hz，强调频率位置来自真实 Hz，而不是普通 token index。之后进入 CoTAR Encoder ×4，沿频率 token 维度 Mean Pooling，得到 frequency feature zf [B,64]。\n\n"
        "在三条分支中均绘制 CoTAR Encoder ×4。不要绘制 Multi-Head Self-Attention。CoTAR 单层放大图应包含：token projection → global core generation → softmax weighting over token dimension → weighted global core aggregation → broadcast global core to all tokens → concatenate [token; global core] → MLP transformation → residual connection and LayerNorm → feed-forward MLP → residual connection and LayerNorm。标注 Branch-wise global core aggregation 或 Centralized global context modeling。\n\n"
        "三条分支的输出分别为 zc、zt、zf，并进入右侧 Branch Fusion。当前采用 concat_mlp：LayerNorm(zc)、LayerNorm(zt)、LayerNorm(zf) → Concatenate [B,192] → Linear 192→512 → GELU → Dropout → Linear 512→256 → Dropout，得到 fused feature [B,256]。当前 concat_mlp 模式不是简单加权求和，不要画成三个分支直接相加。\n\n"
        "最后绘制 Linear Projector：fused feature [B,256] → Linear → class logits [B,40]。主模型图只需保留 logits 输出，不必强制绘制 Softmax。若同时展示第一阶段训练目标，从 fused feature 分出虚线支路：L2 Normalization → embedding z [B,256]，然后分为 CosFace margin head 和 semi-hard triplet mining。CosFace 使用 target logit=s(cosθy−m)，non-target logit=s cosθj；Triplet 分支产生 LTriplet，最终标注 Ltotal=λcls LCosFace+λtri LTriplet。\n\n"
        "版式要求：三条分支上下平行排列，输入位于左侧，融合和输出位于右侧；Channel、Temporal、Frequency 使用蓝色、绿色、橙色/紫色区分；CoTAR 用深色小模块表示，Frequency preprocessing 用连续流程框表示，物理频率编码用带 Hz 标识的独立模块表示；使用实线箭头表示前向传播，虚线箭头表示训练监督；所有张量形状使用 [B,N,D] 标注。整体风格简洁、对齐、无装饰性图片、无人物、无 3D 效果、无复杂背景，适合论文排版。"
    )
    for idx, line in enumerate(prompt.split("\n")):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        set_font(r, name="Microsoft YaHei", size=9.1, color="263238")


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, size=10.5, color="222222")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, color, before, after in [
        (1, 16, BLUE, 16, 8),
        (2, 13, BLUE, 12, 6),
        (3, 11.5, DARK_BLUE, 8, 4),
    ]:
        style = styles[f"Heading {level}"]
        set_style_font(style, size=size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name, size, color in [("Flow", 9.5, NAVY), ("Note", 10.5, "263238"), ("Prompt", 9.1, "263238")]:
        if name not in [s.name for s in styles]:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[name]
        set_style_font(style, name="Microsoft YaHei", size=size, color=color)
        style.paragraph_format.space_before = Pt(3)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.18

    # Quiet running header and footer.
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("DA4FE  |  Model Framework Drawing Brief")
    set_font(hr, size=8.5, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("DA4FE 模型框架图绘制说明  ·  Page ")
    set_font(fr, size=8.5, color=MUTED)
    add_field(fp, "PAGE")

    # Opening masthead.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("DA4FE 模型框架图绘制说明")
    set_font(r, size=25, color=NAVY, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(15)
    r = p.add_run("用于论文级 AI 绘图的软件提示词、结构标注与实现一致性检查")
    set_font(r, size=12.5, color=MUTED)

    add_note(
        doc,
        "核心表达",
        "DA4FE 将 EEG 从通道、时间和频率三个互补视角编码；每条分支通过 CoTAR 提取并广播分支内的全局核心信息，随后进行多分支融合，形成统一的 EEG 表征。",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )

    meta = doc.add_table(rows=4, cols=2)
    set_table_geometry(meta, [1800, 7560])
    set_table_borders(meta, color="E5E7EB", size="4")
    metadata = [
        ("文档用途", "生成用于论文的 DA4FE 模型框架结构图"),
        ("参考实现", "E:\\TeCh-improve\\models\\DA4FE.py；E:\\TeCh-improve\\layers\\Components.py；E:\\TeCh-improve\\layers\\Transformer_EncDec.py"),
        ("默认图示配置", "T=500，C=128，patch_len=4，三分支维度=64，CoTAR 层数=4，融合维度=256，类别数=40"),
        ("建议图形风格", "白底、横向、二维矢量、简洁学术风格；用颜色区分三条分支，用虚线表示训练监督"),
    ]
    for row, (label, value) in zip(meta.rows, metadata):
        set_cell_shading(row.cells[0], "F8FAFC")
        p1 = row.cells[0].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(label)
        set_font(r1, size=9.2, color=DARK_BLUE, bold=True)
        p2 = row.cells[1].paragraphs[0]
        r2 = p2.add_run(value)
        set_font(r2, size=9.2, color="263238")

    add_heading(doc, "1. 模型概述与主线", 1)
    add_body(doc, "DA4FE 的外部输入始终是 X ∈ R[B,T,C]，即 batch、time、channel。模型内部各分支会将输入转置为 [B,C,T]，但图中左侧输入必须标注为 [B,T,C]。三条分支并行工作：Channel Branch 把每个 EEG 通道视为一个 token；Temporal Branch 将时间序列切分为跨通道时间 patch；Frequency Branch 先计算整段信号的单边对数功率谱，再构造带真实物理频率坐标的频率 token。")
    add_body(doc, "三条分支分别经过 CoTAR Encoder 进行分支内全局核心聚合，然后沿各自 token 维度进行 Mean Pooling。三个分支向量经 LayerNorm、拼接和 MLP 融合后，得到最终 fused feature，并连接分类投影层。")
    add_param_table(doc)

    add_heading(doc, "2. 画布布局与视觉组织", 1)
    add_body(doc, "建议采用横向左到右布局。左侧放输入和预处理，中部上下排列三条分支，右侧依次放 Mean Pooling、Branch Fusion、Fused Feature 和分类输出。三条分支必须从同一个输入复制出来，并保持并行，跨分支交互只发生在最后的融合模块。")
    add_flow(doc, "Input [B,T,C]  →  Fixed-length resampling / optional normalization  →  three independent branches  →  branch-wise CoTAR  →  mean pooling  →  fusion  →  classifier", fill=LIGHT_BLUE, accent=NAVY)
    add_body(doc, "配色建议：Channel Branch 使用蓝色，Temporal Branch 使用绿色，Frequency Branch 使用橙色或紫色，CoTAR 使用深蓝色小模块，融合模块使用深灰色或红色。实线箭头表示前向传播，虚线箭头表示第一阶段 CosFace 和 Triplet 训练监督。")

    add_heading(doc, "3. 输入与预处理", 1)
    add_heading(doc, "3.1 输入格式", 2)
    add_body(doc, "左侧输入框写为：Raw EEG，X ∈ R[B,T,C]，Example: [B,500,128]。输入先经过 Fixed-length resampling，将不同长度的 EEG 调整到统一 T；随后可选地进行逐样本、逐通道 z-score normalization。之后将输入复制为三路。")
    add_note(doc, "形状提醒", "图中外部输入是 [B,T,C]；只有进入三个分支后，代码才执行 transpose 得到 [B,C,T]。", fill=LIGHT_GREEN, accent=GREEN)
    add_heading(doc, "3.2 独立增强", 2)
    add_body(doc, "在三个分支入口分别放置 training-time augmentation 小模块。每个分支独立从增强列表中选择一种方式，增强后再进入该分支的 token 构造。不要将三种增强画成一条串联流水线，也不要让增强模块遮挡主数据流。")

    add_heading(doc, "4. Channel Branch：通道分支", 1)
    add_body(doc, "通道分支从 EEG 通道维度建模跨通道关系。每一个 EEG 通道对应一个 token，token 的内容是该通道的完整时间序列。该分支不执行 FFT，也不执行时间 patching。")
    add_flow(doc, "X [B,T,C]  →  Transpose [B,C,T]  →  augmentation  →  sinusoidal positional encoding  →  Linear R^T→R^Dv  →  channel tokens [B,C,Dv]  →  CoTAR Encoder ×4  →  Mean Pooling over channels  →  zc [B,Dv]", fill=LIGHT_BLUE, accent=BLUE, font_size=9.0)
    add_body(doc, "当前配置下，Channel Branch 的形状为 [B,500,128] → [B,128,64] → CoTAR ×4 → [B,64]。图中可标注 Channel-wise Linear Embedding，并注明 Each EEG channel is treated as a token。")

    add_heading(doc, "5. Temporal Branch：时间分支", 1)
    add_body(doc, "时间分支从时间变化中提取局部片段信息。它使用 Cross-Channel Patching，通过 Conv2D 的卷积核同时覆盖全部 EEG 通道和一个局部时间 patch，因此每一个 temporal token 都融合了所有通道在该时间片段中的信息。")
    add_flow(doc, "X [B,T,C]  →  Transpose [B,C,T]  →  augmentation  →  positional encoding  →  Cross-Channel Temporal Patching  →  temporal tokens [B,Nt,Dt]  →  CoTAR Encoder ×4  →  Mean Pooling over time tokens  →  zt [B,Dt]", fill=LIGHT_GREEN, accent=GREEN, font_size=9.0)
    add_body(doc, "Patching 模块建议标注：Conv2D kernel=[C, patch_len]，stride=patch_len，patch_len=4。当前输入 T=500 时，代码中的 replication padding 会使 token 数约为 126，形状可标注为 [B,128,500] → [B,126,64]。")

    add_heading(doc, "6. Frequency Branch：频率分支", 1)
    add_body(doc, "频率分支是 DA4FE 的频域增强模块。它对每个 EEG 通道使用整段固定长度信号进行一次单边 rFFT，得到全局频谱信息；当前实现不是滑动 STFT。频率分支应在图中绘制为一个连续、可读的频谱预处理流程。")
    add_flow(doc, "X [B,T,C]  →  Transpose [B,C,T]  →  augmentation  →  channel-wise mean removal  →  Hann window  →  one-sided rFFT  →  power |FFT(x)|²  →  PSD normalization  →  one-sided correction  →  relative normalization  →  log(PSD+ε)", fill=LIGHT_ORANGE, accent=ORANGE, font_size=9.0)
    add_heading(doc, "6.1 频谱计算细节", 2)
    add_body(doc, "建议在放大框中写出以下简化公式：x0 = x − mean(x)；xw = x0 ⊙ Hann(T)；S = rFFT(xw)；P = |S|²；PSD = P / (fs · Σ Hann²)；对非 DC 和非 Nyquist 的单边频率 bin 乘以 2；relative power = PSD / ΣPSD；最后取 log(relative power + ε)。")
    add_body(doc, "当前参数 fs=1000 Hz、T=500，因此频率分辨率为 fs/T=2 Hz，单边频率 bin 数为 251，物理频率范围为 0-500 Hz。频率数据形状为 [B,128,500] → [B,128,251]。")
    add_heading(doc, "6.2 频率 token 与物理频率坐标", 2)
    add_flow(doc, "Log power spectrum [B,C,251]  →  Cross-Channel Frequency Patching  →  frequency tokens [B,Nf,64]  →  + Physical Frequency Coordinate Encoding  →  CoTAR Encoder ×4  →  Mean Pooling  →  zf [B,64]", fill=LIGHT_ORANGE, accent=ORANGE, font_size=9.0)
    add_body(doc, "当前 patch_len=4 时，频率 token 数约为 63，形状可标注为 [B,128,251] → [B,63,64]。物理频率编码使用每个 token 的真实中心频率 f，而不是 token 序号。坐标为 f/fNyquist 与 log(1+f)/log(1+fNyquist)，经过 Linear(2→64) → GELU → Linear(64→64)，再与频率 token 相加。")
    add_note(doc, "频率创新标注", "请在模块旁明确写出 Physical frequency in Hz / frequency position based on physical frequency, not token index。", fill=LIGHT_ORANGE, accent=ORANGE)

    add_heading(doc, "7. CoTAR Encoder：分支内全局核心建模", 1)
    add_body(doc, "Channel、Temporal 和 Frequency 三条分支均使用 CoTAR Encoder ×4。图中不必展开四个完全相同的层，可以用四个堆叠小方块表示，并在图下方增加一个 CoTAR Layer 放大框。")
    add_flow(doc, "Token X [B,N,D]  →  token projection  →  core projection  →  softmax over N  →  weighted global core [B,1,Dc]  →  broadcast to N tokens  →  concatenate [token; core]  →  MLP  →  residual + LayerNorm  →  FFN  →  residual + LayerNorm", fill=LIGHT_BLUE, accent=NAVY, font_size=8.9)
    add_body(doc, "放大框中要画出：每个 token 先映射为 core representation；在 token 维度上进行 softmax；对所有 token 加权求和得到一个 branch-specific global core；将 global core 广播给所有 token；再将原 token 与 global core 拼接，经 MLP 更新 token。")
    add_note(doc, "重要", "CoTAR 不是标准 Multi-Head Self-Attention。不要画 Q/K/V、多头注意力、Attention Matrix 或 CLS Token。CoTAR 的核心是 Global core extraction + Global core broadcasting + Token refinement。", fill=LIGHT_ORANGE, accent=ORANGE)

    add_heading(doc, "8. Pooling、融合与输出", 1)
    add_heading(doc, "8.1 分支池化", 2)
    add_body(doc, "每个分支的 CoTAR 输出都沿 token 维度执行 Mean Pooling，而不是使用 CLS token。得到 zc、zt、zf 三个分支特征。当前配置为 zc、zt、zf ∈ R[B,64]。")
    add_flow(doc, "Channel tokens [B,128,64]  → mean → zc [B,64]       Temporal tokens [B,126,64] → mean → zt [B,64]       Frequency tokens [B,63,64] → mean → zf [B,64]", fill=LIGHT_GRAY, accent=BLUE, font_size=8.8)
    add_heading(doc, "8.2 Branch Fusion", 2)
    add_body(doc, "当前实现使用 concat_mlp 融合，而不是 add 加权求和。三个分支特征分别经过 LayerNorm，然后拼接为 192 维向量，再经过两层 MLP。")
    add_flow(doc, "zc [B,64] ─LayerNorm─┐\nzt [B,64] ─LayerNorm─┼→ Concatenate [B,192] → Linear 192→512 → GELU → Dropout → Linear 512→256 → Dropout → fused feature [B,256]\nzf [B,64] ─LayerNorm─┘", fill=LIGHT_BLUE, accent=NAVY, font_size=8.8)
    add_note(doc, "融合模式", "只有在 add 模式下才需要绘制 branch projection 和归一化权重加和。当前 concat_mlp 模式中，图中不要把 channel/temporal/frequency 三个权重画成有效的加权求和。", fill=LIGHT_ORANGE, accent=ORANGE)
    add_heading(doc, "8.3 分类输出", 2)
    add_flow(doc, "fused feature [B,256]  →  Linear Projector  →  class logits [B,40]  →  predicted EEG class", fill=LIGHT_GRAY, accent=NAVY, font_size=9.2)

    add_heading(doc, "9. 第一阶段训练目标的可选插图", 1)
    add_body(doc, "如果框架图同时需要说明 Stage-1 feature training，可从 fused feature 分出一条虚线训练支路。主干网络的 fused feature 先进行 L2 normalization 得到 embedding z，再分别接受 CosFace 分类监督和 Triplet 度量监督。")
    add_flow(doc, "fused feature [B,256]  →  L2 normalization  →  embedding z [B,256]\n                                             ├→ CosFace margin head → classification loss\n                                             └→ Semi-hard triplet mining → triplet margin loss\nTotal: L = λcls·LCosFace + λtri·LTriplet", fill=LIGHT_ORANGE, accent=ORANGE, font_size=9.0)
    add_body(doc, "CosFace 模块可标注 target logit=s(cosθy−m)，non-target logit=s cosθj。当前实验配置为 s=30、m=0.4、λcls=1.0、λtri=0.3；这些属于训练配置，建议放在图右下角的 Stage-1 Training Objective 小框中，而不是放大主干结构。")

    add_prompt_block(doc)

    add_heading(doc, "10. 推荐图名与论文图注", 1)
    add_body(doc, "推荐图名：DA4FE: Dual-Axis and Frequency-enhanced EEG Representation Learning Framework。")
    add_body(doc, "推荐图注：DA4FE receives EEG sequences in the form of [B,T,C] and constructs three complementary representations through channel-wise, temporal, and frequency-domain tokenization. Each branch employs CoTAR encoders to extract and broadcast a branch-specific global core representation. The resulting channel, temporal, and frequency features are mean-pooled, normalized, concatenated, and fused by an MLP to obtain the final EEG representation for classification.")

    add_heading(doc, "11. 论文图准确性检查清单", 1)
    checks = [
        "输入格式必须是 [B,T,C]；内部转置后才是 [B,C,T]。",
        "三条分支必须并行，跨分支交互只发生在 Branch Fusion。",
        "Channel Branch 使用通道 token；Temporal Branch 使用跨通道时间 patch；Frequency Branch 使用全段 rFFT。",
        "频率分支必须包含去均值、Hann 加窗、PSD 归一化、单边修正、相对功率和 log 变换。",
        "频率 token 的位置编码必须标注真实物理 Hz，而不是普通 token index。",
        "CoTAR 不要画成 Multi-Head Attention；必须体现 global core 的提取、广播和 token refinement。",
        "三条分支末端使用 Mean Pooling，不使用 CLS Token。",
        "当前 concat_mlp 融合不是简单加权求和。",
        "当前代码没有实际执行 band-pass filter，图中不要添加频带滤波框。",
        "当前频率分支不是 STFT，不要绘制滑动窗口频谱图。",
    ]
    for item in checks:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        r = p.add_run("□  " + item)
        set_font(r, size=10.2, color="263238")

    add_note(doc, "最终建议", "如果 AI 绘图软件无法稳定生成公式和张量形状，优先让它生成模块位置、颜色和箭头，再在 PowerPoint、draw.io、Figma 或 Illustrator 中手动补充公式、尺寸和形状标注。", fill=LIGHT_GREEN, accent=GREEN)

    # Core properties and save.
    props = doc.core_properties
    props.title = "DA4FE 模型框架图绘制说明"
    props.subject = "DA4FE EEG tri-branch model framework drawing prompt"
    props.author = ""
    props.comments = ""
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
